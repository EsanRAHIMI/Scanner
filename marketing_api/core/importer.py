from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .seed import CHANNELS

HEADER_MAP = {
    "snapshot date": "snapshot_date",
    "campaign / ad set": "campaign",
    "objective": "objective",
    "status": "status",
    "daily budget": "daily_budget",
    "monthly budget": "monthly_budget",
    "spend": "spend",
    "impressions": "impressions",
    "reach": "reach",
    "raw platform results": "raw_results",
    "clicks": "clicks",
    "ctr": "ctr",
    "cost / result": "cost_per_result",
    "data quality / notes": "notes",
    "source / report": "source",
    "record type": "record_type",
    "period start": "period_start",
    "period end": "period_end",
    "verified on": "verified_on",
    "record key": "record_key",
}


def _cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and value.startswith("="):
        return None
    return value


def _sheet_rows(ws) -> list[dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell or "").strip().lower() for cell in rows[0]]
    out = []
    for raw in rows[1:]:
        item: dict[str, Any] = {}
        for idx, header in enumerate(headers):
            key = HEADER_MAP.get(header)
            if not key or idx >= len(raw):
                continue
            item[key] = _cell(raw[idx])
        if any(item.values()):
            out.append(item)
    return out


def _num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return None


def load_workbook_rows(path: str) -> dict[str, Any]:
    book = load_workbook(Path(path), data_only=True, read_only=True)
    records = []
    for channel in CHANNELS:
        if channel not in book.sheetnames:
            continue
        for row in _sheet_rows(book[channel]):
            if not row.get("snapshot_date") and not row.get("campaign"):
                continue
            record_type = row.get("record_type") or "Campaign period"
            if channel == "Pinterest" and str(row.get("status") or "").lower() == "draft" and _num(row.get("spend")) is None:
                record_type = "Draft"
            key = row.get("record_key") or f"import:{channel}:{row.get('snapshot_date')}:{row.get('campaign')}:{record_type}:{row.get('period_start')}"
            records.append(
                {
                    "_id": str(key)[:180],
                    "channel": channel,
                    "snapshot_date": row.get("snapshot_date"),
                    "campaign": row.get("campaign"),
                    "objective": row.get("objective"),
                    "status": row.get("status"),
                    "daily_budget": _num(row.get("daily_budget")),
                    "monthly_budget": _num(row.get("monthly_budget")),
                    "spend": _num(row.get("spend")),
                    "impressions": _num(row.get("impressions")),
                    "reach": _num(row.get("reach")),
                    "raw_results": _num(row.get("raw_results")),
                    "clicks": _num(row.get("clicks")),
                    "ctr": _num(row.get("ctr")),
                    "cost_per_result": _num(row.get("cost_per_result")),
                    "notes": row.get("notes"),
                    "source": row.get("source") or "workbook",
                    "record_type": record_type,
                    "period_start": row.get("period_start"),
                    "period_end": row.get("period_end"),
                    "verified_on": row.get("verified_on"),
                }
            )
    book.close()
    return {"records": records}
